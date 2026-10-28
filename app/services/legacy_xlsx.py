"""Bounded converter for the Korean log workbook, shared by upload and initial CLI import.

No formulas or external links are evaluated; uploaded data is never executed.
All nonempty cells and the exact original binary are archived. Plans/examples are not workouts.
"""
import argparse
from io import BytesIO
from zipfile import ZipFile, BadZipFile
import base64
import csv
import hashlib
import json
import re
from datetime import date, datetime, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

from app.sync_schemas import COLUMNS

KST = ZoneInfo("Asia/Seoul")


def cell_value(value):
    return value.isoformat() if isinstance(value, (datetime, date)) else value


MAX_UPLOAD_BYTES = 5 * 1024 * 1024


def seconds_of(value) -> int | None:
    """A duration cell: "43:03", "0:43:03", "43분 3초", a time-of-day cell, a fraction of a day, or minutes."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return round(value * 86400) if 0 < value < 1 else round(value * 60)
    text = str(value).strip()
    text = text.split("T")[-1]  # a time cell exported as 1899-12-30T00:43:03
    if re.fullmatch(r"\d+(?::\d{1,2}){1,2}(?:\.\d+)?", text):
        parts = [float(x) for x in text.split(":")]
        return round(parts[0] * 60 + parts[1]) if len(parts) == 2 else round(parts[0] * 3600 + parts[1] * 60 + parts[2])
    match = re.fullmatch(r"(?:(\d+)\s*시간)?\s*(?:(\d+)\s*분)?\s*(?:(\d+)\s*초)?", text)
    if match and any(match.groups()):
        h, m, s = (int(x or 0) for x in match.groups())
        return h * 3600 + m * 60 + s
    return None


def number_of(value) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        return float(str(value).replace(",", "").strip())
    except ValueError:
        return None


SIDES = {"좌우", "양쪽", "우/좌", "좌/우", "각각"}


def parse_set_text(text: str) -> tuple[float | None, int, str | None] | None:
    """A weight-record cell as (weight, reps, note): "60kg×8", "4kg×15(좌우)", "20kg×5(워밍업)", "맨몸×15",
    "×10" (weight not written), "좌우 합쳐 30회". A drop set ("25kg×5 + 15kg×8") keeps its first part and
    says the rest in the note. None when the cell is not a weight × reps (e.g. "자세 교정")."""
    parts = re.split(r"\s*\+\s*", str(text).strip())
    found = []
    for part in parts:
        match = re.fullmatch(r"(?:(\d+(?:\.\d+)?)\s*kg|맨몸)?\s*[×x]\s*(\d+)\s*(?:\(([^()]*)\))?", part)
        if match:
            found.append((float(match[1]) if match[1] else None, int(match[2]), (match[3] or "").strip()))
            continue
        match = re.fullmatch(r"좌우 합쳐 (\d+)회", part)
        if not match:
            return None
        found.append((None, int(match[1]), ""))
    weight, reps, label = found[0]
    notes = []
    if label in SIDES:
        notes.append("횟수는 한쪽 기준")
    elif label:
        notes.append(label)
    if len(found) > 1:
        notes.append("드롭 세트: 이어서 " + " + ".join(parts[1:]))
    return weight, reps, " · ".join(notes) or None


def daily_target(candidates: list, title: str):
    """Of the same day's workouts (title, item), the one a daily-log row describes: the only one that day, or
    the only one whose title holds the row's routine ("저녁 러닝" → "러닝 · 저녁 러닝"). None when unclear."""
    if len(candidates) == 1:
        return candidates[0][1]
    def squash(text):
        return re.sub(r"\s+|러닝·", "", str(text))
    key = squash(title)
    hits = [item for name, item in candidates if key and (key in squash(name) or squash(name) in key)]
    return hits[0] if len(hits) == 1 else None


RUN_COLUMNS = {"평균심박": "heart_rate", "최대심박": "max_heart_rate", "평균케이던스": "cadence", "최대케이던스": "max_cadence",
               "상승고도(m)": "elevation_gain_m", "칼로리": "calories"}
RUN_LIMITS = {"heart_rate": (20, 250), "max_heart_rate": (20, 250), "cadence": (0, 300), "max_cadence": (0, 300),
              "elevation_gain_m": (0, 10_000), "calories": (0, 20_000), "distance_m": (0, 1_000_000), "duration_seconds": (0, 86400)}


def run_set(fields: dict) -> dict:
    """Set values for one 러닝 기록 row, keyed by its headers. A value that does not read as a number in range
    is left out (the workout memo still has the original)."""
    values = {"distance_m": (lambda km: km and round(km * 1000))(number_of(fields.get("거리(km)"))),
              "duration_seconds": seconds_of(fields.get("운동시간"))}
    for header, name in RUN_COLUMNS.items():
        values[name] = number_of(fields.get(header))
    for name, value in list(values.items()):
        low, high = RUN_LIMITS[name]
        if value is None or not low <= value <= high:
            values.pop(name)
        elif name != "elevation_gain_m" and name != "calories":
            values[name] = round(value)
    return values


def guard_and_open(original: bytes):
    """ZIP/XML-bounded, formula-free open of an untrusted XLSX. The caller must close the book.

    Shared by the fixed-format converter and the AI import reader so both apply the same
    size limits and never evaluate formulas or external links.
    """
    import openpyxl
    if not original or len(original) > MAX_UPLOAD_BYTES:
        raise ValueError("Excel 파일은 5MB 이하이어야 합니다.")
    try:
        with ZipFile(BytesIO(original)) as z:
            files = z.infolist()
            if len(files) > 1000 or sum(x.file_size for x in files) > 20 * 1024 * 1024:
                raise ValueError("압축 해제 크기가 너무 큽니다. 필요한 기록만 포함해 주세요.")
            from defusedxml.ElementTree import fromstring
            cells = 0
            for item in files:
                if not item.filename.startswith("xl/worksheets/") or not item.filename.endswith(".xml"):
                    continue
                root = fromstring(z.read(item))
                for element in root.iter():
                    tag = element.tag.rsplit("}", 1)[-1]
                    if tag == "row" and int(element.get("r", "0")) > 5000:
                        raise ValueError("시트당 5,000행까지만 지원합니다.")
                    if tag == "c":
                        cells += 1
                        match = re.fullmatch(r"([A-Z]+)([0-9]+)", element.get("r", ""))
                        column = 0
                        if not match: raise ValueError("Excel 셀 주소가 올바르지 않습니다.")
                        for letter in match[1]: column = column * 26 + ord(letter) - 64
                        if column > 50 or int(match[2]) > 5000 or cells > 200000:
                            raise ValueError("시트 크기 또는 셀 수가 지원 범위를 초과합니다.")
            if "xl/workbook.xml" not in z.namelist() or any("vbaProject" in x.filename for x in files):
                raise ValueError("매크로 없는 .xlsx 파일만 지원합니다.")
        book = openpyxl.load_workbook(BytesIO(original), data_only=False, read_only=True, keep_links=False)
    except (BadZipFile, KeyError):
        raise ValueError("올바른 .xlsx 파일이 아닙니다.") from None
    for sheet in book:
        if sheet.max_row is None or sheet.max_column is None:
            sheet.calculate_dimension(force=True)
    if any(s.max_row is None or s.max_column is None or s.max_row > 5000 or s.max_column > 50 for s in book):
        book.close()
        raise ValueError("지원 범위는 시트당 5,000행·50열입니다.")
    if sum(s.max_row * s.max_column for s in book) > 200000:
        book.close()
        raise ValueError("Excel 셀 수가 너무 많습니다.")
    return book


def load_cells(original: bytes, max_sheets: int = 30) -> dict[str, list[list]]:
    """Read an arbitrary workbook into {sheet title: nonempty rows of cell values}.

    No fixed structure is required; the AI maps the cells to records. Formulas are read as
    their stored text, never evaluated. Datetimes become ISO strings via cell_value.
    """
    book = guard_and_open(original)
    try:
        if len(book.sheetnames) > max_sheets:
            raise ValueError(f"시트가 너무 많습니다. 최대 {max_sheets}개까지 지원합니다.")
        return {sheet.title: [[cell_value(v) for v in row]
                              for row in sheet.iter_rows(values_only=True) if any(v is not None for v in row)]
                for sheet in book}
    finally:
        book.close()


def build_plan(original: bytes, filename: str) -> dict:
    """Read the fixed 7-sheet Korean workbook without evaluating formulas or external links."""
    book = guard_and_open(original)
    expected = {"12주 루틴", "러닝 기록", "웨이트 기록", "인바디 기록", "대시보드", "식사 기록", "일일 운동 기록"}
    if len(book.sheetnames) != 7 or set(book.sheetnames) != expected:
        book.close()
        raise ValueError("예상한 7개 시트와 다릅니다. 원본 구조를 검토한 후 변환기를 조정하세요.")
    archive = {s.title: [{"row": i, "values": [cell_value(v) for v in row]}
               for i, row in enumerate(s.values, 1) if any(v is not None for v in row)] for s in book}
    try:
        tables, report = convert_book(book, archive)
    finally:
        book.close()
    return {"sha256": hashlib.sha256(original).hexdigest(), "filename": filename,
            "original": base64.b64encode(original).decode(), "archive": archive,
            "tables": tables, "report": report}


def convert_book(book, archive: dict, *, merge_legacy_daily: bool = True) -> tuple[dict, dict]:
    """Convert an already validated workbook; shared with the Sheets value adapter."""
    required_headers = {
        "웨이트 기록": "날짜|루틴|운동|세트1|세트2|세트3|세트4|반복 범위|RIR|총볼륨(kg)|컨디션(1~5)|통증/불편|메모|주차",
        "러닝 기록": "날짜|세션|거리(km)|운동시간|평균페이스|평균심박|최대심박|평균케이던스|최대케이던스|상승고도(m)|칼로리|VO2max|RPE(1~10)|수면(시간)|체중(kg)|메모|주차",
        "인바디 기록": "측정일|체중(kg)|골격근량(kg)|체지방률(%)|체지방량(kg)|BMI|내장지방레벨|기초대사량(kcal)|근육량(kg)|점수|허리둘레(cm)|메모|주차",
        "식사 기록": "일시|구분|음식/제품|섭취량|칼로리(kcal)|탄수화물(g)|당류(g)|단백질(g)|지방(g)|포화지방(g)|나트륨(mg)|콜레스테롤(mg)|메모|주차",
        "일일 운동 기록": "날짜|예정 루틴|실제 수행|상태|사유/메모|주차",
    }
    for name, header in required_headers.items():
        number = 1 if name == "일일 운동 기록" else 3
        expected_header = header.split("|")
        actual = list(next(book[name].iter_rows(min_row=number, max_row=number, values_only=True)))
        if actual[:len(expected_header)] != expected_header or any(x is not None for x in actual[len(expected_header):]):
            book.close()
            raise ValueError(f"{name}: 헤더 순서/이름이 기존 통합기록 양식과 다릅니다.")
        for row in book[name].iter_rows(min_row=number + 1):
            if any(cell.data_type == "f" for cell in row):
                book.close()
                raise ValueError(f"{name}: 기록 시트의 수식은 계산하지 않습니다. 값으로 붙여 넣어 주세요.")
    rows = {key: [] for key in COLUMNS}
    stamp = datetime.now(timezone.utc).isoformat()
    prefix = "legacy_20260926"
    def add(key, identity, **values):
        record = {"external_id": prefix + "_" + identity, "updated_at": stamp, **values}
        rows[key].append(record)
        return record
    def memo(headers, values, source):
        return "원본: " + source + "\n" + "\n".join(str(h) + ": " + str(cell_value(v)) for h,v in zip(headers,values) if h and v is not None)
    groups = {}
    sheet = book["웨이트 기록"]
    headers = list(next(sheet.iter_rows(min_row=3,max_row=3,values_only=True)))
    strength_rows = 0
    complex_sets = 0
    for i, row in enumerate(sheet.iter_rows(min_row=4,values_only=True),4):
        if row[0] == "입력 예시":
            break
        if row[0] is None:
            continue
        if not isinstance(row[0], datetime):
            raise ValueError(f"웨이트 기록 {i}행 날짜 확인 필요")
        day, routine, exercise = row[0].date().isoformat(), row[1], row[2]
        key = (day,routine)
        if key not in groups:
            groups[key] = add("workouts", "w_" + day.replace("-","") + "_" + str(i), workout_date=day, title=routine,
                               body_part=routine, memo="웨이트 기록에서 이관. 시작·종료 시각 미기록.")
        parent = groups[key]
        strength_rows += 1
        for col, raw in enumerate(row[3:7],1):
            if raw is None:
                continue
            text = str(raw)
            parsed = parse_set_text(text)
            weight, reps, label = parsed or (None, None, None)
            if parsed is None:
                complex_sets += 1
            rir = float(row[8]) if row[8] is not None and re.fullmatch(r"\d+(?:\.\d+)?",str(row[8])) else None
            note = memo(headers,row,f"웨이트 기록 {i}행 / 세트{col}")
            if parsed is None:
                note += "\n중량×횟수를 단일 값으로 확정하지 않아 볼륨 합계에서 제외. 원문: " + text
            elif label and label not in text:
                note += "\n원문: " + text + " (" + label + ")"
            add("workout_sets", f"s_{i}_{col}", workout_external_id=parent["external_id"], exercise=exercise,
                set_number=col, weight=weight, reps=reps, completed=True, rir=rir, memo=note)
    for title, key, start in [("러닝 기록","run",4),("일일 운동 기록","daily",2)]:
        sheet=book[title]; headers=list(next(sheet.iter_rows(min_row=start-1,max_row=start-1,values_only=True)))
        for i,row in enumerate(sheet.iter_rows(min_row=start,values_only=True),start):
            if row[0] is None: continue
            if not isinstance(row[0],datetime): raise ValueError(f"{title} {i}행 날짜 확인 필요")
            day=row[0].date().isoformat(); note=memo(headers,row,f"{title} {i}행")
            if merge_legacy_daily and key=="daily":
                # The daily log summarizes a session the weight or running sheet already has: keep it as that
                # session's memo instead of a second, empty workout.
                same_day=[w for w in rows["workouts"] if w["workout_date"]==day and not w["external_id"].startswith(prefix+"_daily_")]
                target=daily_target([(w["title"],w) for w in same_day], str(row[1]))
                if target is not None:
                    target["memo"] += "\n\n" + note
                    continue
            values={"workout_date":day,"title":("걷기" if row[1]=="걷기" else "러닝 · "+str(row[1])) if key=="run" else str(row[1]),"memo":note}
            time_range = re.search(r"(\d{2}:\d{2})~(\d{2}:\d{2})", str(row[2]))
            if key=="daily" and time_range:
                values.update(started_at=day+"T"+time_range[1]+":00+09:00",ended_at=day+"T"+time_range[2]+":00+09:00")
            workout=add("workouts",f"{key}_{i}",**values)
            details=run_set(dict(zip(headers,row))) if key=="run" else {}
            if details:
                add("workout_sets",f"run_set_{i}",workout_external_id=workout["external_id"],exercise="걷기" if row[1]=="걷기" else "러닝",
                    set_number=1,completed=True,memo=None,**details)
    sheet=book["식사 기록"]; headers=list(next(sheet.iter_rows(min_row=3,max_row=3,values_only=True)))
    for i,row in enumerate(sheet.iter_rows(min_row=4,values_only=True),4):
        if row[0] is None: continue
        if not isinstance(row[0],datetime): raise ValueError(f"식사 기록 {i}행 날짜 확인 필요")
        kind=str(row[1]); meal_type="other"
        if "점심" in kind: meal_type="lunch"
        elif "저녁" in kind: meal_type="dinner"
        elif "아침" in kind: meal_type="breakfast"
        elif any(t in kind for t in ["간식","쉐이크","후식"]): meal_type="snack"
        add("meals",f"meal_{i}",eaten_at=row[0].replace(tzinfo=KST).isoformat(),meal_type=meal_type,name=row[2],
            calories=row[4],carbohydrates=row[5],protein=row[7],fat=row[8],memo=memo(headers,row,f"식사 기록 {i}행"))
    sheet=book["인바디 기록"]; headers=list(next(sheet.iter_rows(min_row=3,max_row=3,values_only=True)))
    for i,row in enumerate(sheet.iter_rows(min_row=4,values_only=True),4):
        if row[0] is None: continue
        if not isinstance(row[0],datetime): raise ValueError(f"인바디 기록 {i}행 날짜 확인 필요")
        at=row[0].replace(tzinfo=KST); note=memo(headers,row,f"인바디 기록 {i}행")
        measured_time = re.search(r"(\d{1,2}):(\d{2}) 측정", str(row[11]))
        if measured_time: at=at.replace(hour=int(measured_time[1]),minute=int(measured_time[2]))
        else: note+="\n측정 시각 미상: 날짜 정렬을 위해 00:00 Asia/Seoul로 저장. 실제 시각을 뜻하지 않음."
        add("body_metrics",f"body_{i}",measured_at=at.isoformat(),weight=row[1],skeletal_muscle=row[2],body_fat=row[3],memo=note)
    notes=["루틴·대시보드·입력 예시는 실제 수행 기록으로 집계하지 않고 원본으로 보존합니다.",
           "러닝은 거리·시간·심박·케이던스·상승고도·칼로리를 기록으로 옮기고, 페이스·VO2max·RPE와 추가 영양소, RIR 범위, 인바디 추가 수치는 기록 메모와 원본에서 확인할 수 있습니다.",
           "복합 세트와 중량 미상 세트는 임의의 중량×횟수를 만들지 않습니다. 어시스트 중량과 편측 횟수는 원문 그대로입니다.",
           "측정 시각이 없는 인바디는 정렬을 위한 00:00으로 저장하며 메모에 표시합니다. 수면 시트가 없어 수면 기록은 생성하지 않습니다."]
    report={"counts":{k:len(v) for k,v in rows.items()},"strength_rows":strength_rows,"complex_sets":complex_sets,
            "source_nonempty_rows":{k:len(v) for k,v in archive.items()},"notes":notes,
            "summary":f"운동 {len(rows['workouts'])}회 · 세트 {len(rows['workout_sets'])}개 · 식단 {len(rows['meals'])}건 · 인바디 {len(rows['body_metrics'])}건 이관. 원본 7개 시트 전체 보존."}
    tables={key:[cols]+[[record.get(c) for c in cols] for record in rows[key]] for key,cols in COLUMNS.items()}
    # Fail closed before producing an importable plan.
    from app.services.sync_service import parse_tables
    parse_tables(tables)
    return tables, report


def prepare(path: Path, directory: Path) -> dict:
    plan = build_plan(path.read_bytes(), path.name)
    directory.mkdir(parents=True, exist_ok=True)
    for key, table in plan["tables"].items():
        with (directory / (key + ".csv")).open("w", encoding="utf-8-sig", newline="") as f:
            csv.writer(f).writerows(table)
    (directory / "plan.json").write_text(json.dumps(plan, ensure_ascii=False), encoding="utf-8")
    (directory / "report.json").write_text(json.dumps(plan["report"], ensure_ascii=False, indent=2), encoding="utf-8")
    return plan["report"]


def apply(plan_path: Path, user_email: str) -> dict:
    from sqlalchemy import func, select, text
    from app.database import SessionLocal, initialize_database
    from app.models import ImportArchive, User
    from app.services.sync_service import apply_rows, parse_tables
    plan=json.loads(plan_path.read_text(encoding="utf-8")); original=base64.b64decode(plan["original"],validate=True)
    if hashlib.sha256(original).hexdigest()!=plan["sha256"]: raise ValueError("원본 해시 불일치")
    parsed=parse_tables(plan["tables"])
    initialize_database()
    with SessionLocal() as session:
        session.execute(text("BEGIN IMMEDIATE"))
        try:
            user=session.scalar(select(User).where(func.lower(User.email)==user_email.strip().lower()))
            if user is None: raise ValueError("--user-email에 해당하는 FitLog 사용자가 없습니다.")
            existing=session.scalar(select(ImportArchive).where(ImportArchive.user_id==user.id,ImportArchive.sha256==plan["sha256"]))
            if existing: return {"status":"already_imported","archive_id":existing.id}
            result=apply_rows(session,user.id,parsed)
            archive=ImportArchive(user_id=user.id,sha256=plan["sha256"],filename=plan["filename"],original=original,
                                  payload=json.dumps(plan["archive"],ensure_ascii=False),report=json.dumps(plan["report"],ensure_ascii=False))
            session.add(archive); session.commit()
            return {"status":"ok","archive_id":archive.id,**result}
        except Exception:
            session.rollback(); raise


if __name__ == "__main__":
    parser=argparse.ArgumentParser(description=__doc__)
    group=parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--prepare",type=Path); group.add_argument("--apply",type=Path)
    parser.add_argument("--output",type=Path)
    parser.add_argument("--user-email",help="apply records to this existing FitLog user")
    args=parser.parse_args()
    if args.prepare and not args.output: parser.error("--prepare requires --output")
    if args.apply and not args.user_email: parser.error("--apply requires --user-email")
    print(json.dumps(prepare(args.prepare,args.output) if args.prepare else apply(args.apply,args.user_email),ensure_ascii=False,indent=2))
