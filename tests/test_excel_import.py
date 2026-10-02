import copy
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta
from io import BytesIO
import unittest
from unittest.mock import patch

from support import TEMP
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select, func
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session
from app.database import Base, SessionLocal, engine
from app.main import app
from app.models import Workout, WorkoutSet, Meal, BodyMetric, ExcelPreview, ImportArchive
from app.services.legacy_xlsx import build_plan

HEADERS={
 "웨이트 기록":"날짜|루틴|운동|세트1|세트2|세트3|세트4|반복 범위|RIR|총볼륨(kg)|컨디션(1~5)|통증/불편|메모|주차",
 "러닝 기록":"날짜|세션|거리(km)|운동시간|평균페이스|평균심박|최대심박|평균케이던스|최대케이던스|상승고도(m)|칼로리|VO2max|RPE(1~10)|수면(시간)|체중(kg)|메모|주차",
 "인바디 기록":"측정일|체중(kg)|골격근량(kg)|체지방률(%)|체지방량(kg)|BMI|내장지방레벨|기초대사량(kcal)|근육량(kg)|점수|허리둘레(cm)|메모|주차",
 "식사 기록":"일시|구분|음식/제품|섭취량|칼로리(kcal)|탄수화물(g)|당류(g)|단백질(g)|지방(g)|포화지방(g)|나트륨(mg)|콜레스테롤(mg)|메모|주차",
 "일일 운동 기록":"날짜|예정 루틴|실제 수행|상태|사유/메모|주차",
}
MIME="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def workbook():
    book=Workbook();book.remove(book.active)
    for name in ["12주 루틴","대시보드",*HEADERS]:
        sheet=book.create_sheet(name)
        if name in HEADERS:
            row=1 if name=="일일 운동 기록" else 3
            for col,value in enumerate(HEADERS[name].split("|"),1):sheet.cell(row,col,value)
        else:sheet.cell(1,1,"보존할 원본")
    for c,v in enumerate([datetime(2026,9,25),"하체","스쿼트","60kg×8",None,None,None,"8~10","2",None,None,None,"원문",1],1):book["웨이트 기록"].cell(4,c,v)
    for c,v in enumerate([datetime(2026,9,25,12),"점심","식사","1인분",500,70,5,30,15,3,100,10,"메모",1],1):book["식사 기록"].cell(4,c,v)
    for c,v in enumerate([datetime(2026,9,25),70,32,18,None,None,None,None,None,None,None,"08:30 측정",1],1):book["인바디 기록"].cell(4,c,v)
    return book


def content(book):
    buf=BytesIO();book.save(buf);book.close();return buf.getvalue()


class ExcelTests(unittest.TestCase):
    def setUp(self):
        self.client=TestClient(app);self.client.__enter__()
        with engine.begin() as c:
            for t in reversed(Base.metadata.sorted_tables):c.execute(t.delete())
        self.original=content(workbook())
    def tearDown(self):self.client.__exit__(None,None,None)
    def upload(self,data=None):
        r=self.client.post("/api/sync/excel/preview?filename=records.xlsx",content=self.original if data is None else data,headers={"content-type":MIME})
        self.assertEqual(r.status_code,200,r.text);return r.json()
    def save(self,preview,selected=None):
        if selected is None:selected=[e["key"] for e in preview["entries"] if e["status"] in {"new","updated"}]
        return self.client.post("/api/sync/excel/"+preview["token"]+"/commit",json={"selected":selected})
    def count(self,model):
        with SessionLocal() as s:return s.scalar(select(func.count()).select_from(model))
    def seed(self):
        preview=self.upload();r=self.save(preview);self.assertEqual(r.status_code,200,r.text);return preview
    def test_preview_no_record_write_and_commit_idempotent(self):
        preview=self.upload();self.assertEqual(preview["counts"],{"new":4});self.assertEqual(self.count(Workout),0)
        r=self.save(preview);self.assertEqual(r.status_code,200,r.text);self.assertEqual(r.json()["inserted"],4)
        self.assertEqual(self.save(preview).json(),r.json())
        self.assertEqual(self.count(Workout),1);self.assertEqual(self.count(WorkoutSet),1)
        with SessionLocal() as s:self.assertEqual(s.scalar(select(ImportArchive)).original,self.original)
        again=self.upload();self.assertEqual(again["counts"],{"identical":4})
        self.assertEqual(self.save(again).json()["inserted"],0)
        self.assertEqual(self.count(ImportArchive),1)
    def test_shift_rows_and_reordered_meals_do_not_duplicate(self):
        self.seed();book=load_workbook(BytesIO(self.original))
        for name in ["웨이트 기록","식사 기록","인바디 기록"]:book[name].insert_rows(4,2)
        p=self.upload(content(book));self.assertEqual(p["counts"],{"identical":4})
    def test_select_only_meal_update(self):
        self.seed();book=load_workbook(BytesIO(self.original));book["식사 기록"]["H4"]=42;book["인바디 기록"]["B4"]=71
        p=self.upload(content(book));updates=[e for e in p["entries"] if e["status"]=="updated"]
        self.assertEqual(len(updates),2)
        meal=next(e for e in updates if e["kind"]=="meals");self.assertEqual(meal["before"]["protein"],30);self.assertEqual(meal["incoming"]["protein"],42)
        r=self.save(p,[meal["key"]]);self.assertEqual(r.status_code,200,r.text)
        with SessionLocal() as s:self.assertEqual(s.scalar(select(Meal)).protein,42);self.assertEqual(s.scalar(select(BodyMetric)).weight,70)
        history=self.client.get("/api/sync/excel/history").json();self.assertEqual(history[0]["changes"][0]["before"]["protein"],30)
    def test_missing_rows_are_not_deleted(self):
        self.seed();book=load_workbook(BytesIO(self.original));book["식사 기록"].delete_rows(4)
        p=self.upload(content(book));self.assertEqual(self.save(p).status_code,200);self.assertEqual(self.count(Meal),1)
    def test_duplicate_incoming_is_conflict(self):
        book=load_workbook(BytesIO(self.original));book["식사 기록"].append([c.value for c in book["식사 기록"][4]])
        p=self.upload(content(book));bad=[e for e in p["entries"] if e["status"]=="conflict"]
        self.assertEqual(len(bad),2);self.assertEqual(self.save(p,[bad[0]["key"]]).status_code,422)
        self.assertEqual(self.count(Meal),0)
    def test_multiple_database_candidates_are_conflict(self):
        self.seed()
        with SessionLocal() as s:
            m=s.scalar(select(Meal));s.add(Meal(eaten_at=m.eaten_at,meal_type=m.meal_type,name=m.name));s.commit()
        p=self.upload();self.assertEqual(p["counts"]["conflict"],1)
    def test_stale_preview_rejected(self):
        self.seed();book=load_workbook(BytesIO(self.original));book["식사 기록"]["H4"]=42;p=self.upload(content(book))
        with SessionLocal() as s:m=s.scalar(select(Meal));m.protein=99;s.commit()
        self.assertEqual(self.save(p).status_code,409)
        with SessionLocal() as s:self.assertEqual(s.scalar(select(Meal)).protein,99)
    def test_new_parent_must_be_selected(self):
        p=self.upload();child=next(e for e in p["entries"] if e["kind"]=="workout_sets")
        self.assertEqual(self.save(p,[child["key"]]).status_code,422);self.assertEqual(self.count(Workout),0)
    def test_atomic_failure_rolls_back_all_records(self):
        p=self.upload();original_flush=Session.flush
        def fail(session,*args,**kwargs):
            if any(isinstance(x,Meal) for x in session.new):raise IntegrityError("insert",{},Exception("forced failure"))
            return original_flush(session,*args,**kwargs)
        with patch.object(Session,"flush",fail):r=self.save(p)
        self.assertEqual(r.status_code,409);self.assertEqual(self.count(Workout),0);self.assertEqual(self.count(WorkoutSet),0);self.assertEqual(self.count(ImportArchive),0)
    def test_concurrent_commit_only_once(self):
        p=self.upload()
        with ThreadPoolExecutor(max_workers=2) as pool:results=list(pool.map(lambda _:self.save(p),range(2)))
        self.assertTrue(all(r.status_code==200 for r in results));self.assertEqual(results[0].json(),results[1].json());self.assertEqual(self.count(Meal),1)
    def test_expiry_cancel_and_wrong_selection(self):
        p=self.upload();self.assertEqual(self.save(p,["unknown"]).status_code,422)
        with SessionLocal() as s:s.get(ExcelPreview,p["token"]).created_at-=timedelta(hours=2);s.commit()
        self.assertEqual(self.save(p).status_code,409)
        self.client.delete("/api/sync/excel/"+p["token"]);self.assertEqual(self.save(p).status_code,404)
    def test_archive_only_changes_and_history(self):
        self.seed();book=load_workbook(BytesIO(self.original));book["12주 루틴"]["A1"]="수정한 계획"
        p=self.upload(content(book));self.assertEqual(p["counts"],{"identical":4});self.assertEqual(self.save(p,[]).status_code,200)
        self.assertEqual(self.count(ImportArchive),2);self.assertEqual(self.count(Meal),1)
    def test_existing_google_identity_preserved(self):
        self.seed()
        with SessionLocal() as s:m=s.scalar(select(Meal));m.source="google_sheets";m.external_id="existing_id";s.commit()
        book=load_workbook(BytesIO(self.original));book["식사 기록"]["H4"]=42;p=self.upload(content(book));self.assertEqual(self.save(p).status_code,200)
        with SessionLocal() as s:m=s.scalar(select(Meal));self.assertEqual((m.source,m.external_id),("google_sheets","existing_id"))
    def test_headers_formulas_and_bad_files(self):
        for edit in [lambda b:b["식사 기록"].cell(3,8,"wrong"),lambda b:b["식사 기록"].cell(4,8,"=1+1")]:
            b=workbook();edit(b);r=self.client.post("/api/sync/excel/preview?filename=a.xlsx",content=content(b),headers={"content-type":MIME});self.assertEqual(r.status_code,422,r.text)
        r=self.client.post("/api/sync/excel/preview?filename=a.xlsx",content=b"bad",headers={"content-type":MIME});self.assertEqual(r.status_code,422)
        self.assertEqual(self.count(Workout),0)
    def test_size_bound_and_cross_site(self):
        r=self.client.post("/api/sync/excel/preview?filename=a.xlsx",content=b"x"*(5*1024*1024+1),headers={"content-type":MIME});self.assertEqual(r.status_code,413)
        r=self.client.post("/api/sync/excel/preview?filename=a.xlsx",content=self.original,headers={"content-type":MIME,"sec-fetch-site":"cross-site"});self.assertEqual(r.status_code,403)
    def test_sheet_only_natural_key_changes_show_candidates(self):
        self.seed();b=load_workbook(BytesIO(self.original));b["식사 기록"]["C4"]="다른 이름";p=self.upload(content(b))
        meal=next(e for e in p["entries"] if e["kind"]=="meals");self.assertEqual(meal["status"],"new");self.assertEqual(len(meal["candidates"]),1)

    def test_missing_sheet_dimensions_supported(self):
        import re
        from zipfile import ZipFile
        out=BytesIO()
        with ZipFile(BytesIO(self.original)) as source, ZipFile(out,"w") as target:
            for item in source.infolist():
                data=source.read(item)
                if item.filename.startswith("xl/worksheets/"):
                    data=re.sub(rb"<dimension[^>]*/>",b"",data)
                target.writestr(item,data)
        self.assertEqual(self.upload(out.getvalue())["counts"],{"new":4})

    def test_huge_declared_cell_rejected_before_loading(self):
        from zipfile import ZipFile
        out=BytesIO()
        with ZipFile(BytesIO(self.original)) as source, ZipFile(out,"w") as target:
            for item in source.infolist():
                data=source.read(item)
                if item.filename=="xl/worksheets/sheet1.xml":data=data.replace(b'r="A1"',b'r="A1000000000"')
                target.writestr(item,data)
        r=self.client.post("/api/sync/excel/preview?filename=a.xlsx",content=out.getvalue(),headers={"content-type":MIME})
        self.assertEqual(r.status_code,422)

    def test_fields_absent_from_workbook_are_preserved(self):
        self.seed()
        with SessionLocal() as s:
            s.scalar(select(Meal)).image_url="https://example.com/photo.jpg"
            s.scalar(select(WorkoutSet)).rpe=8
            s.commit()
        p=self.upload();self.assertEqual(p["counts"],{"identical":4})
        b=load_workbook(BytesIO(self.original));b["식사 기록"]["H4"]=40;b["웨이트 기록"]["D4"]="65kg×8"
        p=self.upload(content(b));self.assertEqual(self.save(p).status_code,200)
        with SessionLocal() as s:
            self.assertEqual(s.scalar(select(Meal)).image_url,"https://example.com/photo.jpg")
            self.assertEqual(s.scalar(select(WorkoutSet)).rpe,8)
